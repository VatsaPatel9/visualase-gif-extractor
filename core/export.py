"""Write the extracted data to Excel (live distance formulas) or CSV, honouring the user's choices."""
import csv
import datetime
import io
import os
import zipfile

import numpy as np
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ALL_FIELDS = ['status', 'temps', 'limits', 'positions', 'dist_center', 'dist_tip']

FONT = 'Arial'
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FT = Font(name=FONT, size=14, bold=True)
FBLUE = Font(name=FONT, size=10, color='0000FF')
FNOTE = Font(name=FONT, size=9, italic=True, color='555555')
HDR = PatternFill('solid', start_color='D9E1F2')
GRP = PatternFill('solid', start_color='B4C6E7')
INP = PatternFill('solid', start_color='FFFF00')
THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='center')


def _cell(ws, r, c, v, font=F, fill=None, fmt=None, align=None, border=False):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fill:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align
    if border:
        cell.border = BOX
    return cell


def point_title(key, multi_row):
    row, lab = (int(v) for v in key.split(':'))
    if not multi_row:
        return f'Point {lab}'
    return f'Point {lab} ({"top" if row == 0 else "bottom"} panel)'


def select_frames(job, mode='all', start=None, end=None):
    frames = job.frames
    if mode == 'laser_on':
        return [f for f in frames if f['laser'] == 'On']
    if mode == 'last_per_run':
        last = {}
        for f in frames:
            if f['points'] or f['time']:
                last[f['run']] = f
        return [last[k] for k in sorted(last)]
    if mode == 'range':
        s = 0 if start is None else int(start)
        e = len(frames) - 1 if end is None else int(end)
        return [f for f in frames if s <= f['frame'] <= e]
    return list(frames)


def _fiber_geom(s):
    if not s.get('fiber'):
        return None
    e, t = np.asarray(s['fiber']['entry'], float), np.asarray(s['fiber']['tip'], float)
    return e, t, (e + t) / 2


def point_positions(job, frames, keys):
    """Distinct (slice, point, x, y) positions with the frames they were seen in."""
    pos = {}
    for f in frames:
        for k in keys:
            v = f['points'].get(k)
            if not v or v[2] is None:
                continue
            row = k.split(':')[0]
            sidx = f['slice'].get(row)
            pos.setdefault((sidx, k, v[2], v[3]), []).append(f)
    order = sorted(pos, key=lambda q: (q[0] if q[0] is not None else -1,
                                       int(q[1].split(':')[1]), pos[q][0]['frame']))
    return [(q, pos[q]) for q in order]


def summary_rows(job, frames, keys, pixel_mm):
    out = []
    for (sidx, k, x, y), fl in point_positions(job, frames, keys):
        s = job.slices[sidx] if sidx is not None else None
        g = _fiber_geom(s) if s else None
        dc = dt = None
        if g:
            dc = float(np.hypot(x - g[2][0], y - g[2][1]) * pixel_mm)
            dt = float(np.hypot(x - g[1][0], y - g[1][1]) * pixel_mm)
        lab = int(k.split(':')[1])
        out.append({'slice': s['label'] if s else '', 'slice_idx': sidx, 'key': k, 'point': lab, 'x': x, 'y': y,
                    'first_frame': fl[0]['frame'], 'last_frame': fl[-1]['frame'], 'n_frames': len(fl),
                    'first_time': fl[0]['time'], 'last_time': fl[-1]['time'],
                    'dist_center_mm': None if dc is None else round(dc, 2),
                    'dist_tip_mm': None if dt is None else round(dt, 2),
                    'result': f'{lab}: {dc:.1f} mm' if dc is not None else f'{lab}: (no fiber set)'})
    return out


def export(job, out_dir, fields=None, points=None, frame_mode='all', start=None, end=None,
           fmt='xlsx', pixel_mm=1.015625):
    fields = [f for f in (fields or ALL_FIELDS) if f in ALL_FIELDS]
    keys = [k for k in (points if points is not None else job.point_keys) if k in job.point_keys]
    frames = select_frames(job, frame_mode, start, end)
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(job.filename)[0]
    rows = summary_rows(job, frames, keys, pixel_mm)
    if fmt == 'csv':
        path = _write_csv(job, out_dir, base, fields, keys, frames, rows, pixel_mm)
    else:
        path = _write_xlsx(job, out_dir, base, fields, keys, frames, rows, pixel_mm, frame_mode)
    return path, rows


# ---------------------------------------------------------------- CSV
def _frame_columns(job, fields, keys):
    multi = job.rows > 1
    cols = ['Frame', 'Run', 'Time (mm:ss)', 'Time (s)']
    if 'status' in fields:
        cols += ['Laser', 'Power (%)', 'Watts']
    need_xy = any(f in fields for f in ('positions', 'dist_center', 'dist_tip'))
    if need_xy:
        cols += [f'Slice ({"top" if r == 0 else "bottom"} panel)' if multi else 'Slice' for r in range(job.rows)]
    sub = []
    if 'temps' in fields:
        sub.append('Temp (°C)')
    if 'limits' in fields:
        sub.append('Limit (°C)')
    if need_xy:
        sub += ['X (px)', 'Y (px)']
    if 'dist_center' in fields:
        sub.append('Dist. to fiber center (mm)')
    if 'dist_tip' in fields:
        sub.append('Dist. to fiber tip (mm)')
    return cols, sub, need_xy


def _write_csv(job, out_dir, base, fields, keys, frames, rows, pixel_mm):
    cols, sub, need_xy = _frame_columns(job, fields, keys)
    multi = job.rows > 1
    header = cols + [f'{point_title(k, multi)} {h}' for k in keys for h in sub] + ['Notes']
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for f in frames:
        line = [f['frame'], f['run'], f['time'], f['secs']]
        if 'status' in fields:
            line += [f['laser'], f['power'], f['watts']]
        if need_xy:
            line += [job.slices[f['slice'][str(r)]]['label'] if str(r) in f['slice'] else '' for r in range(job.rows)]
        for k in keys:
            v = f['points'].get(k)
            temp, lim, x, y = v if v else (None, None, None, None)
            g = None
            if v and x is not None:
                s = job.slices[f['slice'][k.split(':')[0]]]
                g = _fiber_geom(s)
            vals = []
            if 'temps' in fields:
                vals.append(temp)
            if 'limits' in fields:
                vals.append(lim)
            if need_xy:
                vals += [x, y]
            if 'dist_center' in fields:
                vals.append(round(float(np.hypot(x - g[2][0], y - g[2][1]) * pixel_mm), 2) if g else None)
            if 'dist_tip' in fields:
                vals.append(round(float(np.hypot(x - g[1][0], y - g[1][1]) * pixel_mm), 2) if g else None)
            line += vals
        line.append('; '.join(f['notes']))
        w.writerow(['' if v is None else v for v in line])
    sbuf = io.StringIO()
    sw = csv.writer(sbuf)
    sw.writerow(['Slice', 'Point', 'X (px)', 'Y (px)', 'First frame', 'Last frame', 'Frames', 'First time', 'Last time',
                 'Distance to fiber center (mm)', 'Distance to fiber tip (mm)', 'Result'])
    for r in rows:
        sw.writerow([r['slice'], r['point'], r['x'], r['y'], r['first_frame'], r['last_frame'], r['n_frames'],
                     r['first_time'], r['last_time'], r['dist_center_mm'], r['dist_tip_mm'], r['result']])
    fbuf = io.StringIO()
    fw = csv.writer(fbuf)
    fw.writerow(['Slice', 'Entry X', 'Entry Y', 'Tip X', 'Tip Y', 'Center X', 'Center Y', 'Pixel size (mm)'])
    for s in job.slices:
        g = _fiber_geom(s)
        if g:
            fw.writerow([s['label'], *np.round(g[0], 2), *np.round(g[1], 2), *np.round(g[2], 2), pixel_mm])
    path = os.path.join(out_dir, f'{base}_csv.zip')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr(f'{base}_frame_data.csv', buf.getvalue())
        z.writestr(f'{base}_distance_summary.csv', sbuf.getvalue())
        z.writestr(f'{base}_fiber.csv', fbuf.getvalue())
    return path


# ---------------------------------------------------------------- Excel
def _write_xlsx(job, out_dir, base, fields, keys, frames, rows, pixel_mm, frame_mode):
    multi = job.rows > 1
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    r = 1
    _cell(ws, r, 1, f'Visualase GIF data extraction - {job.filename}', FT)
    r += 2
    mode_txt = {'all': 'All frames', 'laser_on': 'Laser-on frames only', 'last_per_run': 'Last frame of each run',
                'range': f'Frames {frames[0]["frame"]}-{frames[-1]["frame"]}' if frames else 'Frame range'}[frame_mode]
    for k, v in [('Source GIF', job.filename), ('Frames in GIF', job.n_frames), ('Frames exported', f'{len(frames)} ({mode_txt})'),
                 ('Runs (clock restarts)', max((f['run'] for f in job.frames), default=0)),
                 ('Points exported', ', '.join(point_title(k, multi) for k in keys) or 'none'),
                 ('Extracted on', datetime.date.today().isoformat())]:
        _cell(ws, r, 1, k, FB)
        _cell(ws, r, 2, v)
        r += 1
    r += 1
    _cell(ws, r, 1, 'Settings (editable)', FB)
    r += 1
    _cell(ws, r, 1, 'Pixel size (mm / pixel)', FB)
    c = _cell(ws, r, 2, pixel_mm, FBLUE, INP, '0.000000')
    note = ('Visualase 256x256 panels with a 260 mm field of view = 1.015625 mm/pixel (DICOM PixelSpacing of the '
            'planning "T1 SE" series). Change it here if your acquisition differs; every distance updates.')
    c.comment = Comment(note, 'Visualase extractor')
    _cell(ws, r, 3, note, FNOTE)
    wb.defined_names['PixelSize'] = DefinedName('PixelSize', attr_text=f'Summary!$B${r}')
    r += 2

    # fiber table
    _cell(ws, r, 1, 'Laser fiber per slice (pixels in the 256x256 panel; x right, y down; same on the left and right images)', FB)
    r += 1
    hdr = ['Slice', 'GIF frames', 'Skull entry X', 'Skull entry Y', 'Tip X', 'Tip Y', 'Center X', 'Center Y',
           'Fiber length inside skull (mm)', 'Set by']
    for k, h in enumerate(hdr, 1):
        _cell(ws, r, k, h, FB, HDR, align=WRAP, border=True)
    r += 1
    f0 = r
    src_txt = {'saved': 'reviewed earlier (remembered)', 'auto': 'automatic - not reviewed', 'manual': 'reviewed in app'}
    for s in job.slices:
        g = _fiber_geom(s)
        _cell(ws, r, 1, s['label'], border=True)
        fr = s['frames']
        _cell(ws, r, 2, f'{fr[0]}-{fr[-1]}' if len(fr) > 1 else str(fr[0]), border=True)
        if g:
            for k, v in enumerate([g[0][0], g[0][1], g[1][0], g[1][1]], 3):
                _cell(ws, r, k, round(float(v), 2), FBLUE, fmt='0.0', border=True)
            _cell(ws, r, 7, f'=(C{r}+E{r})/2', fmt='0.0', border=True)
            _cell(ws, r, 8, f'=(D{r}+F{r})/2', fmt='0.0', border=True)
            _cell(ws, r, 9, f'=SQRT((E{r}-C{r})^2+(F{r}-D{r})^2)*PixelSize', fmt='0.0', border=True)
        else:
            for k in range(3, 10):
                _cell(ws, r, k, None, border=True)
        _cell(ws, r, 10, src_txt.get(s.get('fiber_source'), 'not set'), border=True)
        r += 1
    f1 = r - 1
    for nm, col in [('FiberSlice', 'A'), ('FiberCX', 'G'), ('FiberCY', 'H'), ('FiberTX', 'E'), ('FiberTY', 'F')]:
        wb.defined_names[nm] = DefinedName(nm, attr_text=f'Summary!${col}${f0}:${col}${f1}')
    _cell(ws, r, 1, 'Center = midpoint of the fiber segment inside the skull (skull entry -> tip). Blue cells are the '
                    'measured fiber coordinates: edit them and all distances update.', FNOTE)
    r += 2

    want_dist = 'dist_center' in fields or 'dist_tip' in fields
    if want_dist or 'positions' in fields:
        _cell(ws, r, 1, 'Distance of each point from the fiber center (one row per point position; a point appears '
                        'on several rows if it was moved)', FB)
        r += 1
        hdr = ['Slice', 'Point', 'Point X', 'Point Y', 'First frame', 'Last frame', 'Frames', 'First time', 'Last time',
               'Distance to fiber center (px)', 'Distance to fiber center (mm)', 'Distance to fiber tip (mm)', 'Result']
        for k, h in enumerate(hdr, 1):
            _cell(ws, r, k, h, FB, HDR, align=WRAP, border=True)
        r += 1
        if not rows:
            _cell(ws, r, 1, 'No numbered points in the exported frames.', FNOTE)
            r += 1
        for row in rows:
            vals = [row['slice'], row['point'], row['x'], row['y'], row['first_frame'], row['last_frame'],
                    row['n_frames'], row['first_time'], row['last_time']]
            for k, v in enumerate(vals, 1):
                _cell(ws, r, k, v, border=True)
            m = f'MATCH(A{r},FiberSlice,0)'
            if row['dist_center_mm'] is not None:
                _cell(ws, r, 10, f'=SQRT((C{r}-INDEX(FiberCX,{m}))^2+(D{r}-INDEX(FiberCY,{m}))^2)', fmt='0.0', border=True)
                _cell(ws, r, 11, f'=J{r}*PixelSize', fmt='0.0', border=True)
                _cell(ws, r, 12, f'=SQRT((C{r}-INDEX(FiberTX,{m}))^2+(D{r}-INDEX(FiberTY,{m}))^2)*PixelSize', fmt='0.0', border=True)
                _cell(ws, r, 13, f'=B{r}&": "&TEXT(K{r},"0.0")&" mm"', FB, border=True)
            else:
                for k in (10, 11, 12):
                    _cell(ws, r, k, None, border=True)
                _cell(ws, r, 13, row['result'], FB, border=True)
            r += 1
        r += 1
    notes = ['Point position = centre of the magenta point number drawn on the temperature image (about +/-2 px).',
             'Distances are in-plane (2D) on the slice where the point is drawn.']
    glitch = [f['frame'] for f in frames if any('partially drawn' in n for n in f['notes'])]
    if glitch:
        notes.append(f'Frames {", ".join(map(str, glitch[:40]))}{"..." if len(glitch) > 40 else ""} are only partially '
                     'drawn in the GIF itself; unreadable values are left blank.')
    for n in notes:
        _cell(ws, r, 1, 'Note: ' + n, FNOTE)
        r += 1
    for k, wdt in enumerate([18, 12, 11, 11, 11, 11, 10, 10, 14, 16, 16, 14, 14], 1):
        ws.column_dimensions[get_column_letter(k)].width = wdt

    # frame data
    wf = wb.create_sheet('Frame data')
    cols, sub, need_xy = _frame_columns(job, fields, keys)
    for k, h in enumerate(cols, 1):
        _cell(wf, 2, k, h, FB, HDR, align=WRAP, border=True)
    slice_col = {}
    if need_xy:
        n_status = 4 + (3 if 'status' in fields else 0)
        for rr in range(job.rows):
            slice_col[rr] = n_status + 1 + rr
    col = len(cols) + 1
    gcols = {}
    for k in keys:
        if not sub:
            break
        gcols[k] = col
        _cell(wf, 1, col, point_title(k, multi), FB, GRP, align=Alignment(horizontal='center'), border=True)
        if len(sub) > 1:
            wf.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + len(sub) - 1)
        for i, h in enumerate(sub):
            _cell(wf, 2, col + i, h, FB, HDR, align=WRAP, border=True)
        col += len(sub)
    notes_col = col
    _cell(wf, 2, notes_col, 'Notes', FB, HDR, border=True)
    for i, f in enumerate(frames):
        rr = 3 + i
        vals = [f['frame'], f['run'], f['time'], f['secs']]
        if 'status' in fields:
            vals += [f['laser'], f['power'], f['watts']]
        for k, v in enumerate(vals, 1):
            _cell(wf, rr, k, v, fmt='0.0' if k == 7 else None)
        for prow, c in slice_col.items():
            sidx = f['slice'].get(str(prow))
            _cell(wf, rr, c, job.slices[sidx]['label'] if sidx is not None else None)
        for k, c in gcols.items():
            v = f['points'].get(k)
            if not v:
                continue
            temp, lim, x, y = v
            j = c
            if 'temps' in fields:
                _cell(wf, rr, j, temp, fmt='0.0')
                j += 1
            if 'limits' in fields:
                _cell(wf, rr, j, lim)
                j += 1
            if need_xy:
                _cell(wf, rr, j, x)
                _cell(wf, rr, j + 1, y)
                X, Y = get_column_letter(j), get_column_letter(j + 1)
                sc = get_column_letter(slice_col[int(k.split(':')[0])])
                m = f'MATCH({sc}{rr},FiberSlice,0)'
                j += 2
                if 'dist_center' in fields:
                    _cell(wf, rr, j, f'=IF({X}{rr}="","",IFERROR(SQRT(({X}{rr}-INDEX(FiberCX,{m}))^2+'
                                     f'({Y}{rr}-INDEX(FiberCY,{m}))^2)*PixelSize,""))', fmt='0.0')
                    j += 1
                if 'dist_tip' in fields:
                    _cell(wf, rr, j, f'=IF({X}{rr}="","",IFERROR(SQRT(({X}{rr}-INDEX(FiberTX,{m}))^2+'
                                     f'({Y}{rr}-INDEX(FiberTY,{m}))^2)*PixelSize,""))', fmt='0.0')
                    j += 1
        if f['notes']:
            _cell(wf, rr, notes_col, '; '.join(f['notes']), FNOTE)
    for k in range(1, notes_col):
        wf.column_dimensions[get_column_letter(k)].width = 10
    for c in slice_col.values():
        wf.column_dimensions[get_column_letter(c)].width = 15
    wf.column_dimensions[get_column_letter(notes_col)].width = 70
    wf.row_dimensions[2].height = 42
    wf.freeze_panes = wf.cell(row=3, column=5)
    wb.calculation.fullCalcOnLoad = True
    path = os.path.join(out_dir, f'{base}.xlsx')
    wb.save(path)
    return path
